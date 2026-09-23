import os
import sys, json
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

BASE_THEMES = {
    "formal-report": {
        "heading_font": "Times New Roman",
        "body_font": "Times New Roman",
        "heading_sizes": {1: 20, 2: 15, 3: 12},
        "body_size": 11,
        "title_size": 26,
    },
    "internal-memo": {
        "heading_font": "Arial",
        "body_font": "Arial",
        "heading_sizes": {1: 16, 2: 13, 3: 11},
        "body_size": 10,
        "title_size": 20,
    },
    "proposal": {
        "heading_font": "Arial",
        "body_font": "Arial",
        "heading_sizes": {1: 22, 2: 16, 3: 13},
        "body_size": 11,
        "title_size": 20
    },
}

ACCENT_COLORS = {
    "blue":  {"primary_color": RGBColor(0x1F, 0x4E, 0x79), "body_color": RGBColor(0x1A, 0x1A, 0x1A)},
    "green": {"primary_color": RGBColor(0x1E, 0x5E, 0x3A), "body_color": RGBColor(0x1A, 0x1A, 0x1A)},
    "grey":  {"primary_color": RGBColor(0x40, 0x40, 0x40), "body_color": RGBColor(0x1A, 0x1A, 0x1A)},
    "red":   {"primary_color": RGBColor(0x8A, 0x1C, 0x1C), "body_color": RGBColor(0x1A, 0x1A, 0x1A)},
}

ALLOWED_THEME_VARIANTS = [
    "formal-report-blue",
    "formal-report-green",
    "formal-report-grey",
    "internal-memo-grey",
    "internal-memo-blue",
    "proposal-red",
    "proposal-blue",
]

DEFAULT_THEME = "formal-report-blue"

def resolve_theme(theme_name: str)-> dict:
    if theme_name not in ALLOWED_THEME_VARIANTS:
        raise ValueError(
            f"Unknown theme '{theme_name}'. Must be one of: {ALLOWED_THEME_VARIANTS}"
        )

    accent_name = theme_name.rsplit("-",1)[-1]
    base_name = theme_name[: -(len(accent_name)+1)]

    base = BASE_THEMES[base_name]
    accent = ACCENT_COLORS[accent_name]

    return {**base, **accent}

def _apply_run_style(run, font_name: str, size_pt: int, color: RGBColor, bold: bool = False):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.font.color.rgb = color
    run.font.bold = bold

def _add_title(doc: Document, title: str, theme: dict):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    _apply_run_style(run, theme["heading_font"], theme["title_size"], theme["primary_color"], bold=True)
    p.paragraph_format.space_after = Pt(18)

def _add_heading(doc: Document, text: str, level: int, theme: dict):
    level = max(1, min(level, 3))
    p = doc.add_paragraph()
    run = p.add_run(text)
    _apply_run_style(
        run,
        theme["heading_font"],
        theme["heading_sizes"][level],
        theme["primary_color"],
        bold=True,
    )
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)

def _add_paragraph(doc: Document, text: str, theme: dict):
    p = doc.add_paragraph()
    run = p.add_run(text)
    _apply_run_style(run, theme["body_font"], theme["body_size"], theme["body_color"])
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15

def _add_bullet_list(doc: Document, items: list, theme: dict):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(item)
        _apply_run_style(run, theme["body_font"], theme["body_size"], theme["body_color"])

def _add_table(doc: Document, headers: list, rows: list, theme: dict):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Light Grid Accent 1"
 
    header_cells = table.rows[0].cells
    for i, header_text in enumerate(headers):
        header_cells[i].text = ""
        run = header_cells[i].paragraphs[0].add_run(header_text)
        _apply_run_style(run, theme["heading_font"], theme["body_size"], theme["primary_color"], bold=True)

    for row_data in rows:
        row_cells = table.add_row().cells
        for i, cell_text in enumerate(row_data):
            row_cells[i].text = ""
            run = row_cells[i].paragraphs[0].add_run(str(cell_text))
            _apply_run_style(run, theme["body_font"], theme["body_size"], theme["body_color"])
 
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

def build_document(theme_name: str, title: str, sections: list, output_path: str) -> str:
    theme = resolve_theme(theme_name=theme_name)
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    _add_title(doc, title, theme)

    for i, s in enumerate(sections):
        s_type = s.get("type")

        def required(field: str):
            if field not in s:
                raise ValueError(
                    f"Section {i} (type={s_type!r}) is missing required field"
                    f"'{field}', Full section: {s}"
                )
            return s[field]
        
        if s_type == "heading":
            _add_heading(doc, s["text"], s.get("level", 1), theme)
        elif s_type == "paragraph":
            _add_paragraph(doc, s["text"], theme)
        elif s_type == "bullet_list":
            _add_bullet_list(doc, s["items"], theme)
        elif s_type == "table":
            _add_table(doc, s["headers"], s["rows"], theme)
        else:
            raise ValueError(
                f"Section {i}: unknown section type {s_type!r}. "
                f"Must be one of: heading, paragraph, bullet_list, table."
            )

    output_dir = os.path.dirname(output_path)
    if output_dir:
        os.makedirs(output_dir, exist_ok=True)

    doc.save(output_path)
    return output_path

if __name__ == "__main__":
    try:
        raw = sys.stdin.read()
        request = json.loads(raw)

    except json.JSONDecodeError as e:
        print(json.dumps({"success": False, "error": f"Invalid JSON on stdin: {e}"}))
        sys.exit(1)

    try:
        output = build_document(
                theme_name=request["theme"],
                title=request["title"],
                sections=request["sections"],
                output_path=request["output_path"],
            )
        print(json.dumps({"success": True, "output_path": output}))
        sys.exit(0)

    except KeyError as e:
        print(json.dumps({
            "success": False,
            "error": f"Missing required field in request: {e}",
        }))
        sys.exit(1)

    except ValueError as e:
        print(json.dumps({
            "success": False,
            "error": str(e),
        }))
        sys.exit(1)

    except Exception as e:
        print(json.dumps({
            "success": False,
            "error": f"Unexpected error ({type(e).__name__}): {e}",
        }))
        sys.exit(1)